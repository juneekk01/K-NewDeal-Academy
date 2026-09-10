import csv
import hashlib
import io
import json
import math
import uuid
from typing import Dict, List, Tuple

import streamlit as st
from docx import Document
from openai import OpenAI
from pypdf import PdfReader


st.set_page_config(
    page_title="Linear LLM Workflow Builder",
    page_icon="🔗",
    layout="wide",
)


# -----------------------------
# Session state
# -----------------------------
def new_agent(index: int) -> Dict:
    return {
        "id": str(uuid.uuid4()),
        "name": f"Agent {index}",
        "model": "gpt-5.6-luna",
        "system_prompt": "You are a helpful AI agent. Follow the user's request precisely.",
        "handoff_prompt": "",
        "rag_enabled": False,
        "rag_top_k": 4,
        "files": [],
    }


if "agents" not in st.session_state:
    st.session_state.agents = [new_agent(1)]

if "rag_cache" not in st.session_state:
    st.session_state.rag_cache = {}

if "last_run" not in st.session_state:
    st.session_state.last_run = []


# -----------------------------
# File parsing / RAG helpers
# -----------------------------
def read_uploaded_file(name: str, data: bytes) -> str:
    suffix = name.lower().rsplit(".", 1)[-1] if "." in name else ""

    if suffix in {"txt", "md", "py", "json", "yaml", "yml", "log"}:
        return data.decode("utf-8", errors="replace")

    if suffix == "csv":
        text = data.decode("utf-8-sig", errors="replace")
        rows = []
        reader = csv.reader(io.StringIO(text))
        for row in reader:
            rows.append(" | ".join(row))
        return "\n".join(rows)

    if suffix == "pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    if suffix == "docx":
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)

    raise ValueError(
        f"지원하지 않는 파일 형식입니다: {name}. "
        "현재 txt, md, csv, json, yaml/yml, py, log, pdf, docx를 지원합니다."
    )


def chunk_text(text: str, source: str, chunk_size: int = 1400, overlap: int = 220) -> List[Dict]:
    clean = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if not clean:
        return []

    chunks = []
    start = 0
    chunk_no = 1

    while start < len(clean):
        end = min(len(clean), start + chunk_size)

        # Prefer a natural break near the end of a chunk.
        if end < len(clean):
            candidates = [
                clean.rfind("\n\n", start, end),
                clean.rfind("\n", start, end),
                clean.rfind(". ", start, end),
            ]
            best = max(candidates)
            if best > start + chunk_size // 2:
                end = best + 1

        body = clean[start:end].strip()
        if body:
            chunks.append(
                {
                    "source": source,
                    "chunk_no": chunk_no,
                    "text": body,
                }
            )
            chunk_no += 1

        if end >= len(clean):
            break
        start = max(start + 1, end - overlap)

    return chunks


def embedding(client: OpenAI, text: str, model: str = "text-embedding-3-small") -> List[float]:
    response = client.embeddings.create(model=model, input=text)
    return response.data[0].embedding


def embeddings_batch(
    client: OpenAI, texts: List[str], model: str = "text-embedding-3-small"
) -> List[List[float]]:
    response = client.embeddings.create(model=model, input=texts)
    ordered = sorted(response.data, key=lambda x: x.index)
    return [item.embedding for item in ordered]


def cosine_similarity(a: List[float], b: List[float]) -> float:
    dot = sum(x * y for x, y in zip(a, b))
    norm_a = math.sqrt(sum(x * x for x in a))
    norm_b = math.sqrt(sum(y * y for y in b))
    if norm_a == 0 or norm_b == 0:
        return 0.0
    return dot / (norm_a * norm_b)


def rag_index_key(files: List[Dict], embedding_model: str) -> str:
    h = hashlib.sha256()
    h.update(embedding_model.encode())
    for f in files:
        h.update(f["name"].encode("utf-8", errors="ignore"))
        h.update(f["data"])
    return h.hexdigest()


def build_rag_index(
    client: OpenAI,
    files: List[Dict],
    embedding_model: str = "text-embedding-3-small",
) -> List[Dict]:
    key = rag_index_key(files, embedding_model)

    if key in st.session_state.rag_cache:
        return st.session_state.rag_cache[key]

    chunks = []
    for f in files:
        text = read_uploaded_file(f["name"], f["data"])
        chunks.extend(chunk_text(text, f["name"]))

    if not chunks:
        st.session_state.rag_cache[key] = []
        return []

    texts = [c["text"] for c in chunks]
    vectors = []

    # Keep batches modest for predictable Community Cloud memory usage.
    batch_size = 64
    for i in range(0, len(texts), batch_size):
        vectors.extend(
            embeddings_batch(client, texts[i : i + batch_size], embedding_model)
        )

    index = []
    for chunk, vector in zip(chunks, vectors):
        item = dict(chunk)
        item["embedding"] = vector
        index.append(item)

    st.session_state.rag_cache[key] = index
    return index


def retrieve_context(
    client: OpenAI,
    query: str,
    files: List[Dict],
    top_k: int,
    embedding_model: str = "text-embedding-3-small",
) -> Tuple[str, List[Dict]]:
    index = build_rag_index(client, files, embedding_model)
    if not index:
        return "", []

    q_vec = embedding(client, query, embedding_model)
    ranked = sorted(
        (
            (cosine_similarity(q_vec, item["embedding"]), item)
            for item in index
        ),
        key=lambda x: x[0],
        reverse=True,
    )[:top_k]

    hits = []
    blocks = []
    for score, item in ranked:
        hit = {
            "source": item["source"],
            "chunk_no": item["chunk_no"],
            "score": score,
            "text": item["text"],
        }
        hits.append(hit)
        blocks.append(
            f"[Source: {item['source']} / chunk {item['chunk_no']}]\n{item['text']}"
        )

    return "\n\n---\n\n".join(blocks), hits


# -----------------------------
# Workflow serialization
# -----------------------------
def export_workflow() -> str:
    safe_agents = []
    for agent in st.session_state.agents:
        safe_agents.append(
            {
                "name": agent["name"],
                "model": agent["model"],
                "system_prompt": agent["system_prompt"],
                "handoff_prompt": agent["handoff_prompt"],
                "rag_enabled": agent["rag_enabled"],
                "rag_top_k": agent["rag_top_k"],
            }
        )
    return json.dumps({"agents": safe_agents}, ensure_ascii=False, indent=2)


def import_workflow(raw: str) -> None:
    payload = json.loads(raw)
    incoming = payload.get("agents", [])
    if not isinstance(incoming, list) or not incoming:
        raise ValueError("agents 배열이 비어 있거나 올바르지 않습니다.")

    agents = []
    for i, item in enumerate(incoming, start=1):
        agent = new_agent(i)
        agent["name"] = str(item.get("name", f"Agent {i}"))
        agent["model"] = str(item.get("model", "gpt-5.6-luna"))
        agent["system_prompt"] = str(item.get("system_prompt", ""))
        agent["handoff_prompt"] = str(item.get("handoff_prompt", ""))
        agent["rag_enabled"] = bool(item.get("rag_enabled", False))
        agent["rag_top_k"] = int(item.get("rag_top_k", 4))
        agents.append(agent)

    st.session_state.agents = agents
    st.session_state.last_run = []


# -----------------------------
# UI
# -----------------------------
st.title("🔗 Linear LLM Workflow Builder")
st.caption(
    "LLM 에이전트를 만들고 순서를 편집한 뒤, 앞 단계의 output을 다음 단계의 input으로 전달하는 "
    "선형 워크플로를 실행합니다."
)

with st.sidebar:
    st.header("API 설정")
    api_key = st.text_input(
        "OpenAI API Key",
        type="password",
        placeholder="sk-...",
        help="입력한 키는 이 앱의 현재 Streamlit 세션에서만 사용하며 파일로 저장하지 않습니다.",
    )
    st.caption("API 키는 코드/requirements.txt에 포함되지 않습니다.")

    st.divider()
    st.subheader("Workflow")
    if st.button("＋ 에이전트 추가", use_container_width=True):
        st.session_state.agents.append(new_agent(len(st.session_state.agents) + 1))
        st.rerun()

    st.download_button(
        "Workflow JSON 내보내기",
        data=export_workflow(),
        file_name="workflow.json",
        mime="application/json",
        use_container_width=True,
    )

    import_file = st.file_uploader(
        "Workflow JSON 불러오기",
        type=["json"],
        key="workflow_import",
        help="RAG 파일 자체는 JSON에 포함하지 않습니다.",
    )
    if import_file is not None:
        import_hash = hashlib.sha256(import_file.getvalue()).hexdigest()
        if st.session_state.get("last_import_hash") != import_hash:
            try:
                import_workflow(import_file.getvalue().decode("utf-8"))
                st.session_state.last_import_hash = import_hash
                st.success("Workflow를 불러왔습니다.")
                st.rerun()
            except Exception as e:
                st.error(f"Workflow 불러오기 실패: {e}")


# Visual workflow summary
st.subheader("1. Linear workflow 편집")
if st.session_state.agents:
    flow_names = "  →  ".join(
        f"**{i + 1}. {a['name']}**" for i, a in enumerate(st.session_state.agents)
    )
    st.markdown(flow_names)


for i, agent in enumerate(st.session_state.agents):
    aid = agent["id"]

    with st.expander(
        f"{i + 1}. {agent['name']} · {agent['model']}",
        expanded=True,
    ):
        top1, top2, top3, top4 = st.columns([5, 1, 1, 1])

        with top1:
            agent["name"] = st.text_input(
                "에이전트 이름",
                value=agent["name"],
                key=f"name_{aid}",
            )

        with top2:
            st.write("")
            st.write("")
            if st.button("↑", key=f"up_{aid}", disabled=(i == 0), help="위로 이동"):
                st.session_state.agents[i - 1], st.session_state.agents[i] = (
                    st.session_state.agents[i],
                    st.session_state.agents[i - 1],
                )
                st.rerun()

        with top3:
            st.write("")
            st.write("")
            if st.button(
                "↓",
                key=f"down_{aid}",
                disabled=(i == len(st.session_state.agents) - 1),
                help="아래로 이동",
            ):
                st.session_state.agents[i + 1], st.session_state.agents[i] = (
                    st.session_state.agents[i],
                    st.session_state.agents[i + 1],
                )
                st.rerun()

        with top4:
            st.write("")
            st.write("")
            if st.button(
                "🗑",
                key=f"delete_{aid}",
                disabled=(len(st.session_state.agents) == 1),
                help="에이전트 삭제",
            ):
                st.session_state.agents.pop(i)
                st.rerun()

        col_a, col_b = st.columns([1, 2])

        with col_a:
            agent["model"] = st.text_input(
                "Model ID",
                value=agent["model"],
                key=f"model_{aid}",
                help="예: gpt-5.6-luna, gpt-5.6, gpt-5-mini 등 계정에서 사용 가능한 모델 ID",
            )

            agent["rag_enabled"] = st.toggle(
                "RAG 사용",
                value=agent["rag_enabled"],
                key=f"rag_{aid}",
            )

            if agent["rag_enabled"]:
                agent["rag_top_k"] = st.slider(
                    "검색 chunk 수 (Top-K)",
                    min_value=1,
                    max_value=10,
                    value=int(agent["rag_top_k"]),
                    key=f"topk_{aid}",
                )

                uploaded_files = st.file_uploader(
                    "참조 파일 Drag & Drop",
                    type=["txt", "md", "csv", "json", "yaml", "yml", "py", "log", "pdf", "docx"],
                    accept_multiple_files=True,
                    key=f"files_{aid}",
                )
                if uploaded_files:
                    agent["files"] = [
                        {"name": f.name, "data": f.getvalue()}
                        for f in uploaded_files
                    ]
                    st.caption(
                        "참조 파일: " + ", ".join(f["name"] for f in agent["files"])
                    )
                elif agent["files"]:
                    st.caption(
                        "현재 세션 참조 파일: "
                        + ", ".join(f["name"] for f in agent["files"])
                    )
            else:
                agent["files"] = []

        with col_b:
            agent["system_prompt"] = st.text_area(
                "System prompt",
                value=agent["system_prompt"],
                height=180,
                key=f"system_{aid}",
                help="이 에이전트의 역할, 규칙, 출력 형식 등을 지정합니다.",
            )

            handoff_label = (
                "초기 user prompt에 덧붙일 추가 프롬프트"
                if i == 0
                else "이전 Agent output에 덧붙일 추가 프롬프트"
            )
            agent["handoff_prompt"] = st.text_area(
                handoff_label,
                value=agent["handoff_prompt"],
                height=110,
                key=f"handoff_{aid}",
                placeholder=(
                    "예: 아래 입력을 표 형태로 요약하고 위험 요소를 3개 추출해라."
                ),
                help=(
                    "비워두면 전달받은 입력을 그대로 사용합니다. "
                    "입력 본문은 자동으로 함께 전달됩니다."
                ),
            )


st.divider()
st.subheader("2. Workflow 실행")

user_prompt = st.text_area(
    "User prompt",
    height=160,
    placeholder="첫 번째 에이전트에게 보낼 작업 요청을 입력하세요.",
)

run_col, clear_col = st.columns([1, 1])
with run_col:
    run_clicked = st.button(
        "▶ 전체 Workflow 실행",
        type="primary",
        use_container_width=True,
    )
with clear_col:
    if st.button("실행 결과 지우기", use_container_width=True):
        st.session_state.last_run = []
        st.rerun()


if run_clicked:
    if not api_key.strip():
        st.error("OpenAI API Key를 입력해 주세요.")
    elif not user_prompt.strip():
        st.error("User prompt를 입력해 주세요.")
    else:
        client = OpenAI(api_key=api_key.strip())
        previous_output = user_prompt.strip()
        run_results = []
        progress = st.progress(0, text="Workflow 시작")

        try:
            for idx, agent in enumerate(st.session_state.agents):
                progress.progress(
                    idx / len(st.session_state.agents),
                    text=f"{idx + 1}/{len(st.session_state.agents)} · {agent['name']} 실행 중",
                )

                if idx == 0:
                    base_input = user_prompt.strip()
                    source_label = "Initial user prompt"
                else:
                    base_input = previous_output
                    source_label = f"Previous agent output ({st.session_state.agents[idx - 1]['name']})"

                parts = [
                    f"## Input source\n{source_label}",
                    f"## Input\n{base_input}",
                ]

                if agent["handoff_prompt"].strip():
                    parts.append(
                        f"## Additional instruction\n{agent['handoff_prompt'].strip()}"
                    )

                rag_hits = []
                if agent["rag_enabled"]:
                    if not agent["files"]:
                        st.warning(
                            f"{agent['name']}: RAG가 켜져 있지만 참조 파일이 없어 RAG 없이 실행합니다."
                        )
                    else:
                        query_for_rag = (
                            base_input
                            + "\n\n"
                            + agent["handoff_prompt"].strip()
                        )
                        context, rag_hits = retrieve_context(
                            client=client,
                            query=query_for_rag,
                            files=agent["files"],
                            top_k=int(agent["rag_top_k"]),
                        )
                        if context:
                            parts.append(
                                "## Retrieved reference context\n"
                                "아래 내용은 참고 자료에서 검색된 문맥입니다. "
                                "관련된 경우에만 근거로 사용하고, 자료에 없는 내용을 자료에 있다고 말하지 마세요.\n\n"
                                + context
                            )

                final_input = "\n\n".join(parts)

                response = client.responses.create(
                    model=agent["model"].strip(),
                    instructions=agent["system_prompt"].strip() or None,
                    input=final_input,
                )

                output = (response.output_text or "").strip()
                previous_output = output

                run_results.append(
                    {
                        "agent": agent["name"],
                        "model": agent["model"],
                        "input": final_input,
                        "output": output,
                        "rag_hits": rag_hits,
                    }
                )

            st.session_state.last_run = run_results
            progress.progress(1.0, text="Workflow 완료")
            st.success("전체 Linear Workflow 실행이 완료되었습니다.")

        except Exception as e:
            st.session_state.last_run = run_results
            progress.empty()
            st.error(f"실행 중 오류가 발생했습니다: {e}")


if st.session_state.last_run:
    st.divider()
    st.subheader("3. 실행 결과")

    for i, result in enumerate(st.session_state.last_run, start=1):
        with st.expander(
            f"{i}. {result['agent']} · {result['model']}",
            expanded=True,
        ):
            st.markdown("#### Output")
            st.markdown(result["output"] if result["output"] else "_출력 없음_")

            with st.expander("실제 Agent input 보기"):
                st.code(result["input"], language="markdown")

            if result["rag_hits"]:
                with st.expander(f"RAG 검색 근거 {len(result['rag_hits'])}개"):
                    for hit_no, hit in enumerate(result["rag_hits"], start=1):
                        st.markdown(
                            f"**{hit_no}. {hit['source']} · chunk {hit['chunk_no']} "
                            f"· similarity {hit['score']:.3f}**"
                        )
                        st.write(hit["text"])


st.divider()
st.caption(
    "주의: 업로드 파일과 API 키는 영구 저장하지 않습니다. "
    "Streamlit 세션이 종료되면 사라집니다. 민감한 문서는 조직의 데이터 정책을 확인한 뒤 사용하세요."
)
